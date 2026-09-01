#!/usr/bin/env python3
"""
Lee la columna "Eat" del Excel de peso y genera data/dieta.json.

Métrica: días que cumplo el plan de alimentación en la semana natural en curso.
- En cada bloque de mes: Día (1..31) y "Eat": "X" = cumplí, "-"/vacío = no.
- La dieta se sabe al FINAL del día, así que el denominador son los días YA CERRADOS
  de la semana (lunes → ayer). "Hoy" no cuenta hasta que termina.
- Lee TODOS los bloques de mes, así que la semana que cruza el cambio de mes
  (p. ej. lun 31-ago → dom 6-sep) se cuenta bien uniendo agosto y septiembre.
"""

import json
import sys
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

EXCEL_PATH = (
    "/Users/davidrochgarcia/Library/CloudStorage/"
    "GoogleDrive-davidrochgarcia@gmail.com/My Drive/1. dOS/Salud/Peso.xlsx"
)
OUTPUT_PATH = "data/dieta.json"

MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}


def _sin_acentos(texto: str) -> str:
    t = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in t if not unicodedata.combining(c)).strip().lower()


def cumple(v):
    return v is not None and str(v).strip().upper() == "X"


def escanear_bloques(ws):
    bloques = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if not v:
            continue
        mn = MESES.get(_sin_acentos(str(v)).split()[0]) if str(v).strip() else None
        if mn:
            bloques.append({"mes_num": mn, "col_dia": c, "col_eat": c + 2})
    return bloques


def leer_eat_mapa(ws):
    """Mapa {(mes_num, día): valor Eat} recorriendo todos los bloques de mes."""
    bloques = escanear_bloques(ws)
    if not bloques:  # compatibilidad con el formato clásico A/C
        titulo = ws.cell(1, 1).value
        mn = MESES.get(_sin_acentos(str(titulo)).split()[0]) if titulo else None
        bloques = [{"mes_num": mn, "col_dia": 1, "col_eat": 3}]
    mapa = {}
    for b in bloques:
        if not b["mes_num"]:
            continue
        for r in range(3, ws.max_row + 1):
            d = ws.cell(row=r, column=b["col_dia"]).value
            e = ws.cell(row=r, column=b["col_eat"]).value
            if isinstance(d, (int, float)):
                mapa[(b["mes_num"], int(d))] = e
    return mapa


def main():
    excel = sys.argv[1] if len(sys.argv) > 1 else EXCEL_PATH
    out = sys.argv[2] if len(sys.argv) > 2 else OUTPUT_PATH
    hoy = datetime.now()

    wb = openpyxl.load_workbook(excel, data_only=True)
    ws = wb.worksheets[0]
    mapa = leer_eat_mapa(ws)

    lunes = hoy - timedelta(days=hoy.weekday())
    cerrados = hoy.weekday()  # días ya terminados esta semana (lun..ayer); 0 el lunes

    def eat_de(dia_dt):
        return mapa.get((dia_dt.month, dia_dt.day))

    pattern = []
    cumplidos = 0
    for i in range(7):
        dia_dt = lunes + timedelta(days=i)
        if i < cerrados:
            if cumple(eat_de(dia_dt)):
                pattern.append("ok")
                cumplidos += 1
            else:
                pattern.append("fail")
        elif i == cerrados:
            pattern.append("today")
        else:
            pattern.append("future")

    last_start = lunes - timedelta(days=7)
    last_cumplidos = 0
    for i in range(7):
        if cumple(eat_de(last_start + timedelta(days=i))):
            last_cumplidos += 1

    res = {
        "generated_at": hoy.isoformat(timespec="seconds"),
        "week_start": lunes.date().isoformat(),
        "dias_cumplidos": cumplidos,
        "dias_cerrados": cerrados,
        "total_semana": 7,
        "pattern": pattern,
        "last_week": {"cumplidos": last_cumplidos, "total": 7},
    }

    Path(out).parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        json.dump(res, f, ensure_ascii=False, indent=2)

    print(f"Semana del {lunes.date()} · días cerrados: {cerrados}")
    print(f"Dieta: {cumplidos} de {cerrados}  (semana pasada: {last_cumplidos}/7)")
    print(f"Patrón: {pattern}")
    print(f"Escrito: {out}")


if __name__ == "__main__":
    main()
